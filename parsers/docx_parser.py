import io
import re
from typing import List
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class DocxParser:
    """Parser y generador de documentos DOCX con estilo académico profesional."""

    @staticmethod
    def extract_text(file_bytes: bytes) -> str:
        """Extrae el contenido de un documento DOCX conservando encabezados y estructura."""
        doc = Document(io.BytesIO(file_bytes))
        elements = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style.name.startswith("Heading 1"):
                elements.append(f"# {text}")
            elif paragraph.style.name.startswith("Heading 2"):
                elements.append(f"## {text}")
            elif paragraph.style.name.startswith("Heading 3"):
                elements.append(f"### {text}")
            elif paragraph.style.name.startswith("Heading 4"):
                elements.append(f"#### {text}")
            else:
                # Detectar negritas y cursivas por runs
                md_text = DocxParser._paragraph_to_markdown(paragraph)
                elements.append(md_text)

        for table in doc.tables:
            table_md = DocxParser._table_to_markdown(table)
            if table_md:
                elements.append(table_md)

        return "\n\n".join(elements)

    @staticmethod
    def _paragraph_to_markdown(paragraph) -> str:
        """Convierte un párrafo DOCX a markdown preservando negrita/cursiva de los runs."""
        parts = []
        for run in paragraph.runs:
            text = run.text
            if not text:
                continue
            if run.bold and run.italic:
                parts.append(f"***{text}***")
            elif run.bold:
                parts.append(f"**{text}**")
            elif run.italic:
                parts.append(f"*{text}*")
            else:
                parts.append(text)
        result = "".join(parts)
        return result if result else paragraph.text

    @staticmethod
    def _table_to_markdown(table) -> str:
        rows = []
        for row in table.rows:
            row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append("| " + " | ".join(row_data) + " |")

        if not rows:
            return ""

        col_count = len(table.columns)
        separator = "| " + " | ".join(["---"] * col_count) + " |"
        return rows[0] + "\n" + separator + "\n" + "\n".join(rows[1:])

    @classmethod
    def _add_formatted_runs(cls, paragraph, text: str, base_font: str = "Times New Roman", base_size: float = 11):
        """Parsea tokens de negrita, cursiva y código inline y los añade como Runs de docx."""
        # Dividir texto por delimitadores **bold**, *italic*, `code`
        pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)'
        tokens = re.split(pattern, text)

        for token in tokens:
            if not token:
                continue
            if token.startswith("**") and token.endswith("**") and len(token) > 4:
                run = paragraph.add_run(token[2:-2])
                run.font.name = base_font
                run.font.size = Pt(base_size)
                run.bold = True
            elif token.startswith("*") and token.endswith("*") and len(token) > 2:
                run = paragraph.add_run(token[1:-1])
                run.font.name = base_font
                run.font.size = Pt(base_size)
                run.italic = True
            elif token.startswith("`") and token.endswith("`") and len(token) > 2:
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(124, 58, 237)
            else:
                run = paragraph.add_run(token)
                run.font.name = base_font
                run.font.size = Pt(base_size)

    @classmethod
    def _set_paragraph_spacing(cls, paragraph, before: float = 0, after: float = 6, line_spacing: float = 1.15):
        """Configura espaciado de párrafo estilo journal."""
        pf = paragraph.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line_spacing

    @classmethod
    def _add_horizontal_rule(cls, doc):
        """Añade una línea horizontal como separador visual."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'D1D5DB')
        pBdr.append(bottom)
        pPr.append(pBdr)

    @classmethod
    def _set_heading_color(cls, heading, r: int, g: int, b: int):
        """Aplica color a todos los runs de un heading."""
        for run in heading.runs:
            run.font.color.rgb = RGBColor(r, g, b)

    @classmethod
    def create_docx_from_markdown(cls, text: str) -> bytes:
        """Genera un archivo DOCX profesional con tipografía académica (Times New Roman, A4, espaciado journal)."""
        doc = Document()

        # ── Configurar página A4 con márgenes académicos ──
        for section in doc.sections:
            section.page_width = Cm(21.0)
            section.page_height = Cm(29.7)
            section.top_margin = Cm(2.54)
            section.bottom_margin = Cm(2.54)
            section.left_margin = Cm(2.54)
            section.right_margin = Cm(2.54)

        # ── Configurar estilos base del documento ──
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(11)
        font.color.rgb = RGBColor(31, 41, 55)
        pf = style.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.15
        pf.space_after = Pt(6)

        lines = text.split("\n")
        i = 0
        n = len(lines)
        is_first_heading = True

        while i < n:
            raw_line = lines[i]
            line = raw_line.strip()

            if not line:
                i += 1
                continue

            # Ignorar comentarios de página
            if line.startswith("<!-- [PAGE") and line.endswith("-->"):
                i += 1
                continue

            # ── Bloque de código: ``` ... ``` ──
            if line.startswith("```"):
                code_lines = []
                i += 1
                while i < n and not lines[i].strip().startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                if i < n and lines[i].strip().startswith("```"):
                    i += 1
                code_p = doc.add_paragraph()
                code_p.paragraph_format.left_indent = Cm(0.5)
                code_p.paragraph_format.right_indent = Cm(0.5)
                cls._set_paragraph_spacing(code_p, before=6, after=6, line_spacing=1.0)
                run = code_p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)
                # Fondo gris claro via shading
                shading_elm = OxmlElement('w:shd')
                shading_elm.set(qn('w:val'), 'clear')
                shading_elm.set(qn('w:fill'), 'F9FAFB')
                code_p._p.get_or_add_pPr().append(shading_elm)
                continue

            # ── Tablas Markdown ──
            if line.startswith("|") and line.endswith("|"):
                table_lines = []
                while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                    t_line = lines[i].strip()
                    if not re.match(r'^\|[\s\-:]+([\|][\s\-:]+)+\|$', t_line):
                        cells = [c.strip() for c in t_line.strip("|").split("|")]
                        table_lines.append(cells)
                    i += 1

                if table_lines:
                    num_rows = len(table_lines)
                    num_cols = max(len(row) for row in table_lines)
                    docx_table = doc.add_table(rows=num_rows, cols=num_cols)
                    docx_table.style = 'Table Grid'
                    docx_table.alignment = WD_TABLE_ALIGNMENT.CENTER

                    for r_idx, row_data in enumerate(table_lines):
                        for c_idx in range(num_cols):
                            val = row_data[c_idx] if c_idx < len(row_data) else ""
                            cell = docx_table.cell(r_idx, c_idx)
                            cell_p = cell.paragraphs[0]
                            cell_p.text = ""
                            cls._add_formatted_runs(cell_p, val, base_size=9.5)
                            if r_idx == 0:
                                for r in cell_p.runs:
                                    r.bold = True
                                    r.font.color.rgb = RGBColor(255, 255, 255)
                                # Fondo azul oscuro para header
                                shading = OxmlElement('w:shd')
                                shading.set(qn('w:val'), 'clear')
                                shading.set(qn('w:fill'), '1E3A5F')
                                cell._tc.get_or_add_tcPr().append(shading)

                    doc.add_paragraph()  # Separador
                continue

            # ── Líneas horizontales ──
            if re.match(r'^[-*_]{3,}\s*$', line):
                cls._add_horizontal_rule(doc)
                i += 1
                continue

            # ── Encabezados ──
            if line.startswith("#### "):
                h = doc.add_heading(level=4)
                cls._add_formatted_runs(h, line[5:].strip(), base_font="Times New Roman", base_size=11)
                cls._set_heading_color(h, 55, 65, 81)
                cls._set_paragraph_spacing(h, before=8, after=3, line_spacing=1.15)
            elif line.startswith("### "):
                h = doc.add_heading(level=3)
                cls._add_formatted_runs(h, line[4:].strip(), base_font="Times New Roman", base_size=11.5)
                cls._set_heading_color(h, 17, 24, 39)
                cls._set_paragraph_spacing(h, before=10, after=3, line_spacing=1.15)
            elif line.startswith("## "):
                h = doc.add_heading(level=2)
                cls._add_formatted_runs(h, line[3:].strip(), base_font="Times New Roman", base_size=12.5)
                cls._set_heading_color(h, 37, 99, 235)
                cls._set_paragraph_spacing(h, before=12, after=4, line_spacing=1.15)
            elif line.startswith("# "):
                heading_text = line[2:].strip()
                if is_first_heading:
                    # Título principal centrado
                    h = doc.add_heading(level=1)
                    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cls._add_formatted_runs(h, heading_text, base_font="Times New Roman", base_size=16)
                    cls._set_heading_color(h, 17, 24, 39)
                    cls._set_paragraph_spacing(h, before=8, after=12, line_spacing=1.15)
                    cls._add_horizontal_rule(doc)
                    is_first_heading = False
                else:
                    h = doc.add_heading(level=1)
                    cls._add_formatted_runs(h, heading_text, base_font="Times New Roman", base_size=14)
                    cls._set_heading_color(h, 30, 58, 95)
                    cls._set_paragraph_spacing(h, before=16, after=6, line_spacing=1.15)
            # ── Listas con viñeta ──
            elif re.match(r'^[\*\-]\s+', line):
                item_text = re.sub(r'^[\*\-]\s+', '', line)
                p = doc.add_paragraph(style='List Bullet')
                cls._add_formatted_runs(p, item_text)
                cls._set_paragraph_spacing(p, before=0, after=3, line_spacing=1.15)
            # ── Listas numeradas ──
            elif re.match(r'^\d+\.\s+', line):
                item_text = re.sub(r'^\d+\.\s+', '', line)
                p = doc.add_paragraph(style='List Number')
                cls._add_formatted_runs(p, item_text)
                cls._set_paragraph_spacing(p, before=0, after=3, line_spacing=1.15)
            # ── Citas ──
            elif line.startswith("> "):
                quote_lines = [line[2:].strip()]
                while i + 1 < n and lines[i + 1].strip().startswith("> "):
                    i += 1
                    quote_lines.append(lines[i].strip()[2:].strip())
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.0)
                p.paragraph_format.right_indent = Cm(1.0)
                cls._set_paragraph_spacing(p, before=6, after=6, line_spacing=1.15)
                full_quote = " ".join(quote_lines)
                cls._add_formatted_runs(p, full_quote)
                for r in p.runs:
                    r.italic = True
                    r.font.color.rgb = RGBColor(75, 85, 99)
                # Barra lateral izquierda
                pPr = p._p.get_or_add_pPr()
                pBdr = OxmlElement('w:pBdr')
                left = OxmlElement('w:left')
                left.set(qn('w:val'), 'single')
                left.set(qn('w:sz'), '18')
                left.set(qn('w:space'), '8')
                left.set(qn('w:color'), '2563EB')
                pBdr.append(left)
                pPr.append(pBdr)
            # ── Párrafo común ──
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                cls._add_formatted_runs(p, line)
                cls._set_paragraph_spacing(p, before=0, after=6, line_spacing=1.15)

            i += 1

        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()
